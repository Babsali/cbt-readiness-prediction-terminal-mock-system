from flask import Flask, render_template, request, redirect, url_for, flash, Response
from werkzeug.utils import secure_filename
from werkzeug.security import generate_password_hash
from datetime import datetime, timedelta
import json
import sqlite3
import os
import uuid
import numpy as np
import pandas as pd
import csv
from io import StringIO

app = Flask(__name__)
app.secret_key = 'readiness_secret_key_2026'
app.config['MAX_CONTENT_LENGTH'] = 100 * 1024 * 1024
DB_PATH = "database/cbt_system.db"
UPLOAD_FOLDER = "static/uploads"
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

# ===== LAZY LOAD ML MODEL =====
model = None
scaler = None
model_name = "Model"

def load_model():
    global model, scaler, model_name
    if model is None:
        import joblib
        model = joblib.load('model.pkl')
        try:
            scaler = joblib.load('scaler.pkl')
            model_name = model.__class__.__name__
        except:
            model_name = model.__class__.__name__

def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

def save_uploaded_question_image(file_storage, prefix='question'):
    if file_storage is None or not getattr(file_storage, 'filename', None):
        return None
    filename = secure_filename(file_storage.filename)
    if not filename:
        return None
    ext = os.path.splitext(filename)[1].lower()
    if ext not in {'.png', '.jpg', '.jpeg', '.gif', '.webp'}:
        return None
    unique_name = f"{prefix}_{uuid.uuid4().hex}{ext}"
    save_path = os.path.join(UPLOAD_FOLDER, unique_name)
    file_storage.save(save_path)
    return save_path.replace('\\', '/')

def save_bytes_as_temp_upload(file_bytes, filename_hint='upload.png'):
    ext = os.path.splitext(filename_hint)[1].lower()
    if not ext:
        ext = '.png'
    unique_name = f"{os.path.splitext(filename_hint)[0]}_{uuid.uuid4().hex}{ext}"
    save_path = os.path.join(UPLOAD_FOLDER, unique_name)
    with open(save_path, 'wb') as f:
        f.write(file_bytes)
    return save_path.replace('\\', '/')

def ensure_tables():
    os.makedirs("database", exist_ok=True)
    conn = get_db()
    cur = conn.cursor()
    cur.execute("""CREATE TABLE IF NOT EXISTS Users (user_id INTEGER PRIMARY KEY AUTOINCREMENT, username TEXT UNIQUE NOT NULL, full_name TEXT, password_hash TEXT, role TEXT CHECK(role IN ('Student','Teacher','Admin')), created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP)""")
    cur.execute("""CREATE TABLE IF NOT EXISTS Questions (question_id INTEGER PRIMARY KEY AUTOINCREMENT, question_text TEXT NOT NULL, option_a TEXT, option_b TEXT, option_c TEXT, option_d TEXT, correct_answer TEXT CHECK(correct_answer IN ('A','B','C','D')), difficulty TEXT DEFAULT 'Medium', topic TEXT DEFAULT 'General', source_flag TEXT, image_path TEXT, created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP)""")
    cur.execute("""CREATE TABLE IF NOT EXISTS Exams (exam_id INTEGER PRIMARY KEY AUTOINCREMENT, title TEXT NOT NULL, duration_minutes INTEGER DEFAULT 30, total_marks INTEGER DEFAULT 0, created_by INTEGER, created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP, FOREIGN KEY(created_by) REFERENCES Users(user_id))""")
    cur.execute("""CREATE TABLE IF NOT EXISTS Exam_Questions (exam_id INTEGER, question_id INTEGER, PRIMARY KEY(exam_id, question_id), FOREIGN KEY(exam_id) REFERENCES Exams(exam_id), FOREIGN KEY(question_id) REFERENCES Questions(question_id))""")
    cur.execute("""CREATE TABLE IF NOT EXISTS Student_Attempts (attempt_id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER, exam_id INTEGER, score INTEGER, total INTEGER, attempted_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP, FOREIGN KEY(user_id) REFERENCES Users(user_id), FOREIGN KEY(exam_id) REFERENCES Exams(exam_id))""")
    cur.execute("""CREATE TABLE IF NOT EXISTS Predictions (prediction_id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER, attempt_id INTEGER, model_used TEXT, prediction_result TEXT, confidence REAL, predicted_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP, FOREIGN KEY(user_id) REFERENCES Users(user_id), FOREIGN KEY(attempt_id) REFERENCES Student_Attempts(attempt_id))""")
    cur.execute("""CREATE TABLE IF NOT EXISTS Exam_Events (event_id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER, exam_id INTEGER, event_type TEXT, timestamp DATETIME DEFAULT CURRENT_TIMESTAMP, FOREIGN KEY(user_id) REFERENCES Users(user_id))""")
    cur.execute("""CREATE TABLE IF NOT EXISTS ExamAttempt (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        student_name TEXT NOT NULL,
        matric_no TEXT NOT NULL,
        subject TEXT NOT NULL,
        score REAL DEFAULT 0,
        total_questions INTEGER DEFAULT 0,
        percentage REAL DEFAULT 0,
        readiness_score REAL DEFAULT 0,
        readiness_class TEXT DEFAULT 'At-Risk',
        answers_json TEXT,
        features_json TEXT,
        timestamp DATETIME DEFAULT CURRENT_TIMESTAMP
    )""")
    conn.commit()
    conn.close()

ensure_tables()

# ===== OCR FUNCTIONS =====
def extract_text_from_docx(file):
    from docx import Document
    doc = Document(file)
    return '\n'.join([para.text for para in doc.paragraphs if para.text.strip()])

def preprocess_image_for_ocr(img):
    from PIL import Image, ImageEnhance, ImageFilter, ImageOps
    if img.mode not in ('L', 'RGB', 'RGBA', 'CMYK'):
        img = img.convert('RGB')
    width, height = img.size
    if height > width and height / width > 1.6:
        top_crop = max(80, int(height * 0.14))
        bottom_crop = max(180, int(height * 0.24))
        left_crop = int(width * 0.08)
        right_crop = int(width * 0.08)
        img = img.crop((left_crop, top_crop, width - right_crop, height - bottom_crop))
    grayscale = ImageOps.grayscale(img)
    grayscale = ImageOps.autocontrast(grayscale)
    grayscale = grayscale.filter(ImageFilter.MedianFilter(size=3))
    grayscale = grayscale.filter(ImageFilter.SHARPEN)
    enhancer = ImageEnhance.Contrast(grayscale)
    grayscale = enhancer.enhance(4.0)
    grayscale = grayscale.resize((grayscale.width * 2, grayscale.height * 2), Image.LANCZOS)
    arr = np.array(grayscale)
    if arr.size:
        threshold = max(120, int(np.median(arr) * 0.85))
        grayscale = grayscale.point(lambda x: 0 if x < threshold else 255, '1')
    else:
        grayscale = grayscale.point(lambda x: 0 if x < 170 else 255, '1')
    return grayscale

def try_ocr_with_configs(img):
    import pytesseract
    configs = ['--psm 6', '--psm 11', '--psm 4', '--psm 3']
    results = []
    for cfg in configs:
        try:
            text = pytesseract.image_to_string(img, lang='eng', config=cfg)
            cleaned = text.strip()
            if cleaned:
                results.append(cleaned)
        except Exception:
            continue
    return results

def ocr_image_to_text(img):
    processed = preprocess_image_for_ocr(img)
    results = try_ocr_with_configs(processed)
    if not results:
        return ''
    best_result = max(results, key=len)
    return best_result

def clean_ocr_text(text):
    import re
    if text is None:
        return ""
    text = text.replace('\r\n', '\n').replace('\r', '\n')
    text = text.replace('H,0', 'H2O').replace('FORMULK', 'FORMULA')
    text = re.sub(r'(?<![A-Z])([A-D])\s*[\.)]\s*', r'\n\1) ', text, flags=re.I)
    text = re.sub(r'(?<!\w)([A-D])\s*\|\s*', r'\n\1) ', text, flags=re.I)
    text = text.replace('\u00B0', '°').replace('°CIRC', '°C')
    text = text.replace('\\TEXT', 'TEXT')
    text = text.replace('Q-', 'Q ')
    text = re.sub(r'[\u2018\u2019\u201C\u201D]', "'", text)
    text = re.sub(r'\s+', ' ', text)
    text = text.replace(' )', ')').replace(') ', ') ') # <-- FIXED: added closing '
    text = re.sub(r'\s*\n\s*', '\n', text)
    text = re.sub(r'\n{3,}', '\n\n', text)
    return text.strip()

def split_question_blocks(text):
    import re
    if not text:
        return []
    text = text.replace('\r\n', '\n').replace('\r', '\n').strip()
    if not text:
        return []
    lines = [re.sub(r'\s+', ' ', line).strip() for line in text.split('\n') if line.strip()]
    if not lines:
        return []
    blocks = []
    current = []
    last_kind = None
    def is_option(line):
        return bool(re.match(r'^[A-D]\s*[\.)]\s*', line, flags=re.I))
    def is_numbered_q(line):
        return bool(re.match(r'^(?:\d+|Q\s*\d+)\s*[\.)]\s*', line, flags=re.I))
    for line in lines:
        if is_numbered_q(line):
            if current:
                blocks.append(' '.join(current))
            current = [line]
            last_kind = 'question'
            continue
        if is_option(line):
            if not current:
                current = [line]
            elif last_kind == 'option' and len(current) >= 2 and re.search(r'[?]', line):
                blocks.append(' '.join(current))
                current = [line]
            else:
                current.append(line)
            last_kind = 'option'
            continue
        if current and last_kind == 'option' and re.search(r'[?]$', line):
            blocks.append(' '.join(current))
            current = [line]
            last_kind = 'question'
            continue
        if current:
            current.append(line)
            last_kind = 'text'
        else:
            current = [line]
            last_kind = 'question'
    if current:
        blocks.append(' '.join(current))
    if len(blocks) > 1:
        return blocks
    fallback = re.split(r'(?m)(?=(?:^|\n)\s*[A-D]\s*[\.)]\s*)', text)
    cleaned_fallback = [part.strip() for part in fallback if part.strip() and len(part.strip()) > 8]
    if len(cleaned_fallback) > 1:
        return cleaned_fallback
    return blocks or [text]

def parse_text_to_questions(text):
    import re
    questions = []
    blocks = split_question_blocks(text)
    for block in blocks:
        block = clean_ocr_text(block)
        if not block or len(block) < 10:
            continue
        current = {}
        q_match = re.search(r'^(?:[A-D]\s*[\)\.]\s*)?(.+?)(?=(?:\s+[A-D]\s*[\)\.]|\s+[A-D]\s*$|$))', block, re.S)
        if q_match:
            current['question_text'] = q_match.group(1).strip()
        if current.get('question_text'):
            current['question_text'] = re.sub(r'^[A-D]\s*[\)\.]\s*', '', current['question_text'], flags=re.I)
        parts = re.split(r'(?<![A-Z])[A-D][\)]\s*', block)
        raw_opts = [p.strip() for p in parts[1:] if p.strip()]
        opts = [re.split(r'ANSWER[:\s]*[A-D]', o, flags=re.I)[0].strip() for o in raw_opts]
        if len(opts) >= 2:
            current['option_a'] = opts[0]
            current['option_b'] = opts[1]
            current['option_c'] = opts[2] if len(opts) > 2 else ''
            current['option_d'] = opts[3] if len(opts) > 3 else ''
        ans_match = re.search(r'(?:ANSWER|ANS)[:\s]*([A-D])|(?:^|\s)([A-D])\s*$', block, re.I)
        current['correct_answer'] = (ans_match.group(1) or ans_match.group(2)).upper() if ans_match else 'A'
        if 'question_text' in current and current.get('option_a') and current.get('option_b'):
            current['question_text'] = re.sub(r'\s+', ' ', current['question_text']).strip()
            current['option_a'] = re.sub(r'\s+', ' ', current['option_a']).strip()
            current['option_b'] = re.sub(r'\s+', ' ', current['option_b']).strip()
            current['option_c'] = re.sub(r'\s+', ' ', current.get('option_c', '')).strip()
            current['option_d'] = re.sub(r'\s+', ' ', current.get('option_d', '')).strip()
            current['difficulty'] = rate_difficulty(current['question_text'])
            current['topic'] = 'General'
            questions.append(current)
    return questions

def normalize_csv_columns(df):
    mapping = {
        'question_text': 'question_text', 'Question': 'question_text',
        'option_a': 'option_a', 'A': 'option_a',
        'option_b': 'option_b', 'B': 'option_b',
        'option_c': 'option_c', 'C': 'option_c',
        'option_d': 'option_d', 'D': 'option_d',
        'correct_answer': 'correct_answer', 'Answer': 'correct_answer',
        'difficulty': 'difficulty', 'topic': 'topic'
    }
    df.columns = [mapping.get(col.strip(), col.strip()) for col in df.columns]
    return df

def rate_difficulty(text):
    word_count = len(text.split())
    if word_count < 15: return "Easy"
    elif word_count < 30: return "Medium"
    else: return "Hard"

def smart_parse_from_file(file, ext):
    import pytesseract
    from PIL import Image
    from pdf2image import convert_from_bytes
    pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'
    os.environ["PATH"] += os.pathsep + r'C:\poppler-26.02.0\Library\bin'
    file.stream.seek(0)
    text = ""
    try:
        if ext == 'pdf':
            images = convert_from_bytes(file.read(), poppler_path=r'C:\poppler-26.02.0\Library\bin', dpi=300)
            for img in images:
                text += ocr_image_to_text(img) + "\n"
        elif ext == 'docx':
            text = extract_text_from_docx(file)
        else:
            image = Image.open(file.stream)
            text = ocr_image_to_text(image)
    except Exception as e:
        print("PARSE ERROR:", e)
    return parse_text_to_questions(text)

@app.route("/")
def home():
    return render_template('home.html')

@app.route('/admin/dashboard', methods=['GET'])
def admin_dashboard():
    conn = get_db()
    conn.row_factory = sqlite3.Row

    # Get filter from URL?subject=Math&readiness=Ready
    subject_filter = request.args.get('subject', '')
    readiness_filter = request.args.get('readiness', '')

    user_count = conn.execute('SELECT COUNT(DISTINCT matric_no) as c FROM ExamAttempt').fetchone()['c']
    attempt_count = conn.execute('SELECT COUNT(*) as c FROM ExamAttempt').fetchone()['c']
    exam_count = attempt_count

    # Questions count
    q_count = 0
    for table_name in ['QuestionBank', 'questions', 'Question', 'question_bank']:
        try:
            q_count = conn.execute(f'SELECT COUNT(*) as c FROM {table_name}').fetchone()['c']
            break
        except: continue

    # NEW: Average Readiness Score
    avg_score = conn.execute('SELECT AVG(readiness_score) as avg FROM ExamAttempt').fetchone()['avg']
    if avg_score is None: avg_score = 0.0

    # Build query for table with filters
    query = 'SELECT * FROM ExamAttempt WHERE 1=1'
    params = []
    if subject_filter:
        query += ' AND subject =?'
        params.append(subject_filter)
    if readiness_filter:
        query += ' AND readiness_class =?'
        params.append(readiness_filter)
    query += ' ORDER BY id DESC LIMIT 20'

    attempts = conn.execute(query, params).fetchall()

    # For dropdown: get all unique subjects
    subjects = conn.execute('SELECT DISTINCT subject FROM ExamAttempt ORDER BY subject').fetchall()

    # Data for doughnut chart
    ready = conn.execute("SELECT COUNT(*) FROM ExamAttempt WHERE readiness_class='Ready'").fetchone()[0]
    atrisk = conn.execute("SELECT COUNT(*) FROM ExamAttempt WHERE readiness_class='At-Risk'").fetchone()[0]
    labels = ['Ready', 'At-Risk']
    values = [ready, atrisk]

    conn.close()
    return render_template('admin_dashboard.html',
                           user_count=user_count, q_count=q_count, exam_count=exam_count,
                           attempt_count=attempt_count, avg_score=avg_score, attempts=attempts,
                           subjects=subjects, subject_filter=subject_filter, readiness_filter=readiness_filter,
                           labels=labels, values=values)

@app.route('/admin/download_template')
def download_template():
    df = pd.DataFrame([{
        'question_text': 'What is the capital of Nigeria?',
        'option_a': 'Lagos', 'option_b': 'Abuja', 'option_c': 'Kano', 'option_d': 'Ibadan',
        'correct_answer': 'B', 'difficulty': 'Easy', 'topic': 'Geography'
    }])
    csv_data = df.to_csv(index=False, encoding='utf-8')
    return Response(csv_data, mimetype="text/csv", headers={"Content-disposition": "attachment; filename=question_template.csv"})

@app.route('/admin/upload_questions', methods=["GET","POST"])
def upload_questions():
    if request.method == 'POST':
        file = request.files['file']
        if file.filename == '':
            flash('No file selected', 'danger')
            return redirect(request.url)
        ext = file.filename.rsplit('.', 1)[1].lower()
        questions = []
        image_url = None
        try:
            if ext in ['csv', 'xlsx', 'xls']:
                df = pd.read_csv(file) if ext == 'csv' else pd.read_excel(file)
                questions = normalize_csv_columns(df).to_dict('records')
                flash(f'Loaded {len(questions)} questions from {ext.upper()}', 'success')
            elif ext in ['pdf', 'jpg', 'jpeg', 'png', 'docx']:
                if ext in ['jpg', 'jpeg', 'png']:
                    file_bytes = file.read()
                    saved_path = save_bytes_as_temp_upload(file_bytes, filename_hint=file.filename)
                    image_url = '/' + saved_path
                    class _FS:
                        def __init__(self, path, fname):
                            self.stream = open(path, 'rb')
                            self.filename = fname
                    file_wrapper = _FS(saved_path, file.filename)
                    questions = smart_parse_from_file(file_wrapper, ext)
                    try: file_wrapper.stream.close()
                    except: pass
                else:
                    questions = smart_parse_from_file(file, ext)
                flash(f'Parsed {len(questions)} questions from scanned file. Please review.', 'warning')
            else:
                flash(f'Unsupported file type:.{ext}', 'danger')
                return redirect(request.url)
        except Exception as e:
            flash(f'Error parsing file: {str(e)}', 'danger')
            return redirect(request.url)
        return render_template('review_upload.html', questions=questions, image_url=image_url)
    return render_template('upload_questions.html')

@app.route('/admin/bulk_upload', methods=['GET', 'POST'])
def bulk_upload():
    if request.method == 'POST':
        manual_text = request.form.get('manual_text')
        files = request.files.getlist('bulk_file')
        all_questions = []
        if manual_text: all_questions.extend(parse_text_to_questions(manual_text))
        for file in files:
            if file.filename == '': continue
            ext = file.filename.rsplit('.', 1)[1].lower()
            try:
                if ext in ['csv', 'xlsx', 'xls']:
                    df = pd.read_csv(file) if ext == 'csv' else pd.read_excel(file)
                    all_questions.extend(normalize_csv_columns(df).to_dict('records'))
                elif ext in ['pdf', 'jpg', 'jpeg', 'png', 'docx']:
                    all_questions.extend(smart_parse_from_file(file, ext))
            except Exception as e: flash(f"Error: {file.filename}: {str(e)}", 'danger')
        flash(f'Loaded {len(all_questions)} questions. Please review.', 'success' if all_questions else 'danger')
        return render_template('review_upload.html', questions=all_questions)
    return render_template('bulk_upload.html')

@app.route('/admin/save_reviewed_questions', methods=['POST'])
def save_reviewed_questions():
    conn = get_db()
    count = 0
    i = 0
    while f'q_text_{i}' in request.form:
        image_path = None
        uploaded_image = request.files.get(f'q_image_{i}')
        if uploaded_image and uploaded_image.filename:
            image_path = save_uploaded_question_image(uploaded_image, prefix=f'question_{i}')
        # FIXED: 10 columns = 10?
        conn.execute("""INSERT INTO Questions
        (question_text, option_a, option_b, option_c, option_d, correct_answer, difficulty, topic, source_flag, image_path)
        VALUES (?,?,?,?,?,?,?,?,?,?)""",
        (request.form[f'q_text_{i}'], request.form[f'q_a_{i}'], request.form[f'q_b_{i}'],
         request.form[f'q_c_{i}'], request.form[f'q_d_{i}'], request.form[f'q_ans_{i}'],
         request.form[f'q_diff_{i}'], request.form[f'q_topic_{i}'], 'BULK_UPLOAD', image_path))
        count += 1
        i += 1
    conn.commit()
    conn.close()
    flash(f'Success! {count} questions saved to question bank.', 'success')
    return redirect(url_for('questions_list'))

@app.route("/add_user", methods=["GET","POST"])
def add_user():
    if request.method == "POST":
        conn = get_db()
        try:
            conn.execute("INSERT INTO Users (username, full_name, password_hash, role) VALUES (?,?,?,?)",
                         (request.form['username'], request.form['full_name'], generate_password_hash(request.form['password']), request.form['role']))
            conn.commit()
            flash("User added!", 'success')
            return redirect("/users")
        except Exception as e: flash(f"Error: {e}", 'danger')
        finally: conn.close()
    return render_template("add_user.html")

@app.route("/users")
def users_list():
    conn = get_db()
    users = conn.execute("SELECT * FROM Users ORDER BY user_id DESC").fetchall()
    conn.close()
    return render_template("users.html", users=users)

@app.route('/questions')
@app.route("/questions/search")
def questions_list():
    db = get_db()
    q = request.args.get("q", ""); difficulty = request.args.get("difficulty", ""); topic = request.args.get("topic", "")
    sql = "SELECT * FROM Questions WHERE 1=1"; params = []
    if q: sql += " AND question_text LIKE?"; params.append(f"%{q}%")
    if difficulty: sql += " AND difficulty =?"; params.append(difficulty)
    if topic: sql += " AND topic LIKE?"; params.append(f"%{topic}%")
    sql += " ORDER BY question_id DESC"
    qs = db.execute(sql, params).fetchall()
    db.close()
    return render_template("question_bank.html", questions=qs, q=q, difficulty=difficulty, topic=topic)

@app.route('/questions/edit/<int:question_id>', methods=['GET', 'POST'])
def edit_question(question_id):
    db = get_db()
    question = db.execute("SELECT * FROM Questions WHERE question_id =?", (question_id,)).fetchone()
    if request.method == 'POST':
        image_path = question['image_path'] if question else None
        uploaded_image = request.files.get('image')
        if uploaded_image and uploaded_image.filename:
            image_path = save_uploaded_question_image(uploaded_image, prefix=f'question_edit_{question_id}')
        elif request.form.get('remove_image') == '1':
            image_path = None
        db.execute("""UPDATE Questions SET question_text=?, option_a=?, option_b=?, option_c=?, option_d=?, correct_answer=?, difficulty=?, topic=?, image_path=? WHERE question_id=?""",
            (request.form['question_text'], request.form['option_a'], request.form['option_b'],
             request.form['option_c'], request.form['option_d'], request.form['correct_answer'],
             request.form['difficulty'], request.form['topic'], image_path, question_id))
        db.commit(); db.close()
        flash('Question updated successfully', 'success')
        return redirect(url_for('questions_list'))
    db.close()
    return render_template('edit_question.html', q=question)

@app.route('/questions/delete/<int:question_id>')
def delete_question(question_id):
    db = get_db()
    db.execute("DELETE FROM Questions WHERE question_id =?", (question_id,))
    db.commit(); db.close()
    flash("Question deleted", "success")
    return redirect(url_for("questions_list"))

@app.route('/questions/bulk_delete', methods=['POST'])
def bulk_delete_questions():
    question_ids = request.form.getlist('question_ids')
    if not question_ids:
        flash('No questions selected', 'warning')
        return redirect(url_for('questions_list'))
    db = get_db()
    placeholders = ','.join('?' for _ in question_ids)
    db.execute(f"DELETE FROM Questions WHERE question_id IN ({placeholders})", question_ids)
    db.commit(); db.close()
    flash(f'{len(question_ids)} question(s) deleted successfully', 'success')
    return redirect(url_for('questions_list'))

@app.route('/predict_readiness', methods=['GET', 'POST'])
def predict_readiness():
    load_model()
    prediction_result = None
    if request.method == 'POST':
        try:
            score = float(request.form['score']); total = float(request.form['total']); time_spent = float(request.form['time_spent'])
            accuracy = (score / total) * 100 if total > 0 else 0
            features = np.array([[score, total, time_spent, accuracy]])
            if scaler: features = scaler.transform(features)
            prediction = model.predict(features)[0]; proba = model.predict_proba(features)[0]
            prediction_result = {'class': 'Ready' if prediction == 1 else 'At-Risk', 'confidence': round(max(proba) * 100, 2)}
        except Exception as e: flash(f"Error in prediction: {e}", 'danger')
    return render_template('predict_readiness.html', prediction=prediction_result, model_name=model_name)

@app.route('/exam/login', methods=['GET'])
def exam_login():
    conn = get_db()
    subjects = conn.execute("SELECT DISTINCT topic FROM Questions WHERE topic!= ''").fetchall()
    conn.close()
    subjects = [s['topic'] for s in subjects] if subjects else ['General']
    return render_template('student_login.html', subjects=subjects)

@app.route('/exam/start', methods=['POST'])
def start_exam():
    student_name = request.form['student_name']
    matric_no = request.form['matric_no']
    subject = request.form['subject']
    conn = get_db()
    # FIXED: fetchall not fetchone
    questions = conn.execute("SELECT * FROM Questions WHERE topic =? ORDER BY RANDOM() LIMIT 20", (subject,)).fetchall()
    conn.close()
    if len(questions) < 5:
        flash(f'Not enough questions for {subject}. Add more questions first.', 'danger')
        return redirect(url_for('exam_login'))
    return render_template('take_exam.html', student_name=student_name, matric_no=matric_no, subject=subject, questions=questions, duration=30)

@app.route('/exam/submit', methods=['POST'])
def submit_exam():
    load_model()
    student_name = request.form['student_name']
    matric_no = request.form['matric_no']
    subject = request.form['subject']
    time_spent = float(request.form.get('time_spent', 30*60))
    question_changes = json.loads(request.form.get('question_changes', '{}'))

    answers = {}; score = 0; total = 0
    topic_scores = {} # for consistency

    conn = get_db()
    easy_total=medium_total=hard_total=0
    easy_correct=medium_correct=hard_correct=0

    for key in request.form:
        if key.startswith('q_'):
            q_id = key.split('_')[1]
            student_answer = request.form[key]
            answers[q_id] = student_answer
            q = conn.execute("SELECT correct_answer, difficulty, topic FROM Questions WHERE question_id =?", (q_id,)).fetchone()
            if q:
                total += 1
                if student_answer == q['correct_answer']: score += 1

                # Accuracy by difficulty
                if q['difficulty'] == 'Easy': easy_total += 1
                if q['difficulty'] == 'Medium': medium_total += 1
                if q['difficulty'] == 'Hard': hard_total += 1
                if student_answer == q['correct_answer']:
                    if q['difficulty'] == 'Easy': easy_correct += 1
                    if q['difficulty'] == 'Medium': medium_correct += 1
                    if q['difficulty'] == 'Hard': hard_correct += 1

                # For consistency by topic
                if q['topic'] not in topic_scores: topic_scores[q['topic']] = []
                topic_scores[q['topic']].append(1 if student_answer == q['correct_answer'] else 0)

    percentage = (score / total) * 100 if total > 0 else 0

    # ===== CALCULATE THE 8 THESIS FEATURES =====
    avg_response_time_sec = time_spent / total if total > 0 else 0
    revision_rate = sum(1 for v in question_changes.values() if v > 1) / total if total > 0 else 0
    accuracy_easy = (easy_correct/easy_total)*100 if easy_total>0 else 0
    accuracy_medium = (medium_correct/medium_total)*100 if medium_total>0 else 0
    accuracy_hard = (hard_correct/hard_total)*100 if hard_total>0 else 0
    time_management_score = time_spent / (30*60) # ratio of time used
    completion_rate = total / 20 # assuming 20 questions per exam
    # Consistency = std dev of topic accuracies
    topic_accuracies = [sum(v)/len(v)*100 for v in topic_scores.values() if v]
    consistency_score = float(np.std(topic_accuracies)) if topic_accuracies else 0

    features = np.array([[avg_response_time_sec, revision_rate, accuracy_easy, accuracy_medium, accuracy_hard, time_management_score, completion_rate, consistency_score]])
    if scaler: features = scaler.transform(features)
    proba_array = model.predict_proba(features)[0]

    if percentage >= 50.0:
        readiness_class = 'Ready'; readiness_score = round(proba_array[1] * 100, 2)
    else:
        readiness_class = 'At-Risk'; readiness_score = round(proba_array[0] * 100, 2)

    features_dict = {
        'avg_response_time_sec': avg_response_time_sec, 'revision_rate': revision_rate,
        'accuracy_easy': accuracy_easy, 'accuracy_medium': accuracy_medium, 'accuracy_hard': accuracy_hard,
        'time_management_score': time_management_score, 'completion_rate': completion_rate, 'consistency_score': consistency_score
    }

    conn.execute("""INSERT INTO ExamAttempt
    (student_name, matric_no, subject, score, total_questions, percentage, 
     readiness_score, readiness_class, answers_json, features_json)
    VALUES (?,?,?,?,?,?,?,?,?,?)""",
    (student_name, matric_no, subject, score, total, percentage, 
     readiness_score, readiness_class, json.dumps(answers), json.dumps(features_dict)))
    conn.commit()
    print("SAVED TO DB:", student_name, features_dict)
    conn.close()

    return render_template('exam_result.html', student_name=student_name, score=score, total=total, percentage=round(percentage, 2), readiness_class=readiness_class, readiness_score=readiness_score)

@app.route('/admin/download_csv')
def download_csv():
    conn = get_db()
    conn.row_factory = sqlite3.Row
    attempts = conn.execute('SELECT * FROM ExamAttempt ORDER BY id DESC').fetchall()
    conn.close()
    
    output = StringIO()
    writer = csv.writer(output)
    
    header = [
        'attempt_id', 'student_name', 'matric_no', 'subject', 'score', 'total_questions', 
        'percentage', 'readiness_score', 'readiness_class', 'timestamp',
        'f1_accuracy_easy', 'f2_accuracy_medium', 'f3_accuracy_hard', 
        'f4_avg_response_time_sec', 'f5_completion_rate', 'f6_consistency_score', 
        'f7_revision_rate', 'f8_time_management_score'
    ]
    writer.writerow(header)
    
    for a in attempts:
        features = {}
        if a['features_json']:
            try:
                features = json.loads(a['features_json'])
            except:
                features = {}
        
        row = [
            a['id'], a['student_name'], a['matric_no'], a['subject'], a['score'], a['total_questions'],
            a['percentage'], a['readiness_score'], a['readiness_class'], a['timestamp'],
            features.get('accuracy_easy', ''),
            features.get('accuracy_medium', ''),
            features.get('accuracy_hard', ''),
            features.get('avg_response_time_sec', ''),
            features.get('completion_rate', ''),
            features.get('consistency_score', ''),
            features.get('revision_rate', ''),
            features.get('time_management_score', '')
        ]
        writer.writerow(row)
    
    output.seek(0)
    return Response(output, mimetype="text/csv",
                    headers={"Content-Disposition":"attachment;filename=cbt_attempts_with_features.csv"})

# Auto-create table on startup
with app.app_context():
    conn = sqlite3.connect('database.db')
    conn.execute('''CREATE TABLE IF NOT EXISTS ExamAttempt
        (id INTEGER PRIMARY KEY AUTOINCREMENT,
         student_name TEXT, matric_no TEXT, subject TEXT, score INTEGER, 
         total_questions INTEGER, percentage REAL, readiness_score REAL,
         answers_json TEXT, features_json TEXT, timestamp DATETIME DEFAULT CURRENT_TIMESTAMP,
         readiness_class TEXT)''')
    conn.commit()
    conn.close()
    print("Database table checked/created")

if __name__ == "__main__":
    app.run(debug=False, host='0.0.0.0', port=5000)